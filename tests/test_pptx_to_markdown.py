"""converters/pptx_to_markdown.py 테스트.

python-pptx로 PPTX를 즉석에서 만들어 변환기에 넣고, 슬라이드 제목/본문 텍스트/표/
노트가 Markdown으로 올바르게 옮겨지는지 확인한다. 외부 네트워크나 LLM 호출 없이
전부 로컬에서 동작해야 하므로, 이 테스트도 별도 서버/모델 없이 항상 실행 가능해야
한다.
"""
import subprocess
import sys
from pathlib import Path

import pytest

pptx = pytest.importorskip("pptx")

from pptx.enum.shapes import MSO_SHAPE  # noqa: E402
from pptx.util import Inches  # noqa: E402

from converters.pptx_to_markdown import convert_pptx_file, convert_pptx_to_markdown  # noqa: E402


def _build_sample_pptx(path: Path) -> Path:
    prs = pptx.Presentation()

    # Slide 1: 제목 + 본문 텍스트(불릿)
    layout = prs.slide_layouts[1]  # Title and Content
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = "전극 검사기 사양"
    body = slide.placeholders[1]
    tf = body.text_frame
    tf.text = "폭 800mm 이상"
    p2 = tf.add_paragraph()
    p2.text = "정확도 ±1um"
    p2.level = 1

    notes_slide = slide.notes_slide
    notes_slide.notes_text_frame.text = "이 슬라이드는 초안입니다."

    # Slide 2: 표(사양 데이터)
    blank_layout = prs.slide_layouts[6]
    slide2 = prs.slides.add_slide(blank_layout)
    rows, cols = 2, 2
    table_shape = slide2.shapes.add_table(rows, cols, Inches(1), Inches(1), Inches(4), Inches(1))
    table = table_shape.table
    table.cell(0, 0).text = "항목"
    table.cell(0, 1).text = "값"
    table.cell(1, 0).text = "정확도"
    table.cell(1, 1).text = "±1um"

    prs.save(str(path))
    return path


@pytest.fixture()
def sample_pptx(tmp_path: Path) -> Path:
    return _build_sample_pptx(tmp_path / "sample.pptx")


def test_slide_title_becomes_heading(sample_pptx: Path):
    md = convert_pptx_to_markdown(sample_pptx, extract_images=False)
    assert "## Slide 1: 전극 검사기 사양" in md


def test_body_text_becomes_bullet_list(sample_pptx: Path):
    md = convert_pptx_to_markdown(sample_pptx, extract_images=False)
    assert "- 폭 800mm 이상" in md
    assert "  - 정확도 ±1um" in md  # level=1은 들여쓰기됨


def test_notes_are_included(sample_pptx: Path):
    md = convert_pptx_to_markdown(sample_pptx, extract_images=False)
    assert "**Notes:**" in md
    assert "> 이 슬라이드는 초안입니다." in md


def test_table_becomes_markdown_table(sample_pptx: Path):
    md = convert_pptx_to_markdown(sample_pptx, extract_images=False)
    assert "| 항목 | 값 |" in md
    assert "|---|---|" in md
    assert "| 정확도 | ±1um |" in md


def test_title_is_not_duplicated_in_body(sample_pptx: Path):
    """slide.shapes.title은 호출마다 새 래퍼 객체를 반환하므로, 이전에는 identity(`is`)
    비교가 항상 실패해 제목 도형이 본문에도 한 번 더 불릿으로 찍히는 회귀가 있었다."""
    md = convert_pptx_to_markdown(sample_pptx, extract_images=False)
    assert md.count("전극 검사기 사양") == 1


def test_slide_order_is_preserved(sample_pptx: Path):
    md = convert_pptx_to_markdown(sample_pptx, extract_images=False)
    slide1_idx = md.index("## Slide 1")
    slide2_idx = md.index("## Slide 2")
    assert slide1_idx < slide2_idx


def test_convert_pptx_file_writes_output(tmp_path: Path, sample_pptx: Path):
    out_path = tmp_path / "converted.md"
    result_path = convert_pptx_file(sample_pptx, out_path, extract_images=False)
    assert result_path == out_path
    assert out_path.exists()
    content = out_path.read_text(encoding="utf-8")
    assert "전극 검사기 사양" in content


def test_convert_pptx_file_default_output_path(sample_pptx: Path):
    result_path = convert_pptx_file(sample_pptx, extract_images=False)
    assert result_path == sample_pptx.with_suffix(".md")
    assert result_path.exists()


def test_no_network_or_external_api_import():
    """외부 API 클라이언트(openai 등)를 import하지 않는지 소스 레벨로 확인."""
    source = Path("converters/pptx_to_markdown.py").read_text(encoding="utf-8")
    for forbidden in ("openai", "requests", "httpx", "urllib.request", "socket"):
        assert forbidden not in source, f"외부 통신 관련 모듈({forbidden})을 사용하면 안 됨"


def _build_merged_category_pptx(path: Path) -> Path:
    """실제 사내 사양서 슬라이드에서 흔한 패턴: '구분' 열에 각 행마다 같은 값을
    타이핑한 뒤(H/W x5, Vision x5) 세로로 병합한 표. python-pptx/PowerPoint는
    이렇게 병합하면 병합된 셀의 텍스트가 "H/W\\nH/W\\n..."처럼 줄바꿈으로
    이어붙는다 — 변환기가 이를 중복 없이 "H/W" 한 번으로 정리해야 한다."""
    prs = pptx.Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rows = [
        ("구분", "항목", "내용"),
        ("H/W", "검사 Stage", "세라믹 흡입"),
        ("H/W", "비전시스템 정확도", "0.8um"),
        ("Vision", "카메라 Type", "2D Area"),
        ("Vision", "카메라 해상도", "152M"),
    ]
    table_shape = slide.shapes.add_table(len(rows), 3, Inches(1), Inches(1), Inches(6), Inches(3))
    table = table_shape.table
    for r, (a, b, c) in enumerate(rows):
        table.cell(r, 0).text = a
        table.cell(r, 1).text = b
        table.cell(r, 2).text = c
    table.cell(1, 0).merge(table.cell(2, 0))  # H/W 2행 병합
    table.cell(3, 0).merge(table.cell(4, 0))  # Vision 2행 병합
    prs.save(str(path))
    return path


def test_merged_category_column_fills_every_row(tmp_path: Path):
    pptx_path = _build_merged_category_pptx(tmp_path / "merged.pptx")
    md = convert_pptx_to_markdown(pptx_path, extract_images=False)
    assert "| H/W | 검사 Stage | 세라믹 흡입 |" in md
    assert "| H/W | 비전시스템 정확도 | 0.8um |" in md
    assert "| Vision | 카메라 Type | 2D Area |" in md
    assert "| Vision | 카메라 해상도 | 152M |" in md
    # 병합 과정에서 이어붙은 "H/W\nH/W"가 그대로 노출되면 안 됨
    assert "H/W<br>H/W" not in md
    assert "Vision<br>Vision" not in md


def _build_narrative_plus_table_pptx(path: Path) -> Path:
    """제목 + 도면 placeholder 사각형(AutoShape) + 설명 문단 + 표 하나로 이뤄진
    슬라이드, 그리고 표가 아예 없는 두 번째 슬라이드. tables_only 옵션이 표 없는
    슬라이드는 통째로 생략하고, 표 있는 슬라이드에서는 사각형/문단은 건너뛰고
    표만 남기는지 확인하기 위한 fixture."""
    prs = pptx.Presentation()

    slide1 = prs.slides.add_slide(prs.slide_layouts[6])
    slide1.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(5), Inches(0.5)).text_frame.text = "검사기 H/W"
    box = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(1), Inches(3), Inches(1))
    box.text_frame.text = "검사기 그림"
    para_box = slide1.shapes.add_textbox(Inches(0.5), Inches(2.2), Inches(6), Inches(0.5))
    para_box.text_frame.text = "양극 절연을 검사하여 공정 품질 관리를 위함"
    table_shape = slide1.shapes.add_table(2, 2, Inches(0.5), Inches(3), Inches(4), Inches(1))
    table_shape.table.cell(0, 0).text = "항목"
    table_shape.table.cell(0, 1).text = "내용"
    table_shape.table.cell(1, 0).text = "정확도"
    table_shape.table.cell(1, 1).text = "0.8um"

    slide2 = prs.slides.add_slide(prs.slide_layouts[6])
    slide2.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(5), Inches(0.5)).text_frame.text = "표 없는 슬라이드"
    slide2.shapes.add_textbox(Inches(0.5), Inches(1), Inches(5), Inches(0.5)).text_frame.text = "이 문단은 표와 무관함"

    prs.save(str(path))
    return path


def test_tables_only_skips_narrative_shapes_but_keeps_table(tmp_path: Path):
    pptx_path = _build_narrative_plus_table_pptx(tmp_path / "narrative.pptx")
    md = convert_pptx_to_markdown(pptx_path, extract_images=False, tables_only=True)
    assert "| 항목 | 내용 |" in md
    assert "| 정확도 | 0.8um |" in md
    assert "검사기 그림" not in md
    assert "양극 절연을 검사하여" not in md


def test_tables_only_omits_slides_without_tables(tmp_path: Path):
    pptx_path = _build_narrative_plus_table_pptx(tmp_path / "narrative2.pptx")
    md = convert_pptx_to_markdown(pptx_path, extract_images=False, tables_only=True)
    assert "표 없는 슬라이드" not in md
    assert "이 문단은 표와 무관함" not in md


def test_full_mode_still_includes_narrative_shapes(tmp_path: Path):
    """tables_only=False(기본값)면 기존처럼 도형/문단 텍스트도 그대로 포함되어야 함."""
    pptx_path = _build_narrative_plus_table_pptx(tmp_path / "narrative3.pptx")
    md = convert_pptx_to_markdown(pptx_path, extract_images=False, tables_only=False)
    assert "검사기 그림" in md
    assert "양극 절연을 검사하여" in md
    assert "표 없는 슬라이드" in md


def _build_freeform_title_pptx(path: Path) -> Path:
    """사내 PPT처럼 정식 Title placeholder 없이, 자유 배치 텍스트 상자로 제목을
    넣은 슬라이드(blank 레이아웃). 제목 상자가 맨 위(top이 가장 작음), 그 아래
    요약 설명 한 줄, 더 아래 표 하나가 있다."""
    prs = pptx.Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank: title placeholder 없음
    assert slide.shapes.title is None

    title_box = slide.shapes.add_textbox(Inches(0.4), Inches(0.15), Inches(9), Inches(0.5))
    title_box.text_frame.text = "설비 구성 및 세부 사양"

    summary_box = slide.shapes.add_textbox(Inches(0.4), Inches(0.65), Inches(9), Inches(0.4))
    summary_box.text_frame.text = "요약 설명 한 줄"

    table_shape = slide.shapes.add_table(2, 2, Inches(0.4), Inches(1.2), Inches(4), Inches(1))
    table_shape.table.cell(0, 0).text = "항목"
    table_shape.table.cell(0, 1).text = "사양"
    table_shape.table.cell(1, 0).text = "분해능"
    table_shape.table.cell(1, 1).text = "0.8um"

    prs.save(str(path))
    return path


def test_infers_title_from_topmost_textbox_when_no_placeholder(tmp_path: Path):
    pptx_path = _build_freeform_title_pptx(tmp_path / "freeform_title.pptx")
    md = convert_pptx_to_markdown(pptx_path, extract_images=False)
    assert "## Slide 1: 설비 구성 및 세부 사양" in md
    # 제목으로 쓰인 텍스트가 본문에 불릿으로 중복되면 안 됨
    assert md.count("설비 구성 및 세부 사양") == 1
    # 제목이 아닌 다른 텍스트 상자는 여전히 본문에 남아야 함
    assert "- 요약 설명 한 줄" in md


def test_real_title_placeholder_takes_priority_over_inference(sample_pptx: Path):
    """정식 Title placeholder가 있으면(기존 sample_pptx fixture) 추정 로직을 타지
    않고 placeholder 값을 그대로 써야 한다 — 회귀 방지."""
    md = convert_pptx_to_markdown(sample_pptx, extract_images=False)
    assert "## Slide 1: 전극 검사기 사양" in md


def test_runs_standalone_as_a_script(tmp_path: Path, sample_pptx: Path):
    """`python converters/pptx_to_markdown.py ...`로 이 파일 하나만 직접 실행해도
    (main.py/cli_commands.py를 거치지 않고, FastAPI/dotenv 등 웹 서버 의존성
    없이) 변환이 되는지 확인한다."""
    out_path = tmp_path / "standalone.md"
    module_path = Path(__file__).resolve().parent.parent / "converters" / "pptx_to_markdown.py"

    result = subprocess.run(
        [sys.executable, str(module_path), str(sample_pptx), "-o", str(out_path), "--no-images"],
        capture_output=True,
        text=True,
        timeout=30,
    )

    assert result.returncode == 0, result.stderr
    assert out_path.exists()
    assert "전극 검사기 사양" in out_path.read_text(encoding="utf-8")
